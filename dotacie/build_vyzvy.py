"""Porovnanie dotačných výziev ČR pre grantové poradenstvo -> vyzvy_CZ_porovnanie_2026-09-25.xlsx (pip install openpyxl)."""
import datetime as dt
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from vyzvy_data import CONT, EXP, L, OPEN, R, TODAY, W  # noqa: E402

REL_RANK = {"Vysoká": 0, "Stredná": 1, "Nízka": 2}


def sort_key(item):
    idx, row = item
    date = dt.date.fromisoformat(row[9]) if row[9] else None
    expired = date is not None and date < TODAY
    return (expired, REL_RANK[row[12]], date or dt.date(2100, 1, 1), idx)


rows = [row for _, row in sorted(enumerate(R, 1), key=sort_key)]
ids = [idx for idx, _ in sorted(enumerate(R, 1), key=sort_key)]

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
BODY = Font(name=ARIAL, size=10)
BLUE = Font(name=ARIAL, size=10, color="0000FF")
LINK = Font(name=ARIAL, size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")

wb = Workbook()
summary = wb.active
summary.title = "Súhrn"
calls = wb.create_sheet("Výzvy")
legend = wb.create_sheet("Legenda")
DNES = "'Súhrn'!$B$3"

# --- Výzvy ---------------------------------------------------------------------------------
headers = ["ID", "Zdroj", "Poskytovateľ / program", "Výzva", "Pre koho", "Typ financovania", "Územie",
           "Alokácia", "Max. %", "Max. dotácia", "Uzávierka", "Termín – poznámka", "Status v zdroji",
           "Stav k dátumu", "Dní do uzávierky", "Relevancia pre poradenstvo", "Poznámka", "Overenie / zdroj"]
widths = [5, 13, 24, 48, 30, 16, 20, 18, 11, 18, 12, 30, 13, 17, 11, 14, 44, 30]
for col, (name, width) in enumerate(zip(headers, widths), 1):
    cell = calls.cell(row=1, column=col, value=name)
    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    calls.column_dimensions[get_column_letter(col)].width = width
calls.row_dimensions[1].height = 32

for r, (idx, row) in enumerate(zip(ids, rows), 2):
    (src, prov, name, who, kind, area, alloc, pct, maxg, deadline, note_t, status, rel, note, verify) = row
    values = [idx, src, prov, name, who, kind, area, alloc, pct, maxg,
              dt.date.fromisoformat(deadline) if deadline else None, note_t or None, status]
    for col, value in enumerate(values, 1):
        cell = calls.cell(row=r, column=col, value=value)
        cell.font, cell.border, cell.alignment = BODY, BORDER, WRAP
    calls.cell(row=r, column=11).number_format = "d. m. yyyy"
    k, m = f"K{r}", f"M{r}"
    calls.cell(row=r, column=14, value=(
        f'=IF({k}="",{m},IF({k}<{DNES},"Po uzávierke",IF({m}="{EXP}","{EXP}",'
        f'IF({k}-{DNES}<=14,"Končí do 14 dní","{OPEN}"))))'))
    calls.cell(row=r, column=15, value=f'=IF({k}="","",{k}-{DNES})')
    calls.cell(row=r, column=15).number_format = "0"
    for col, value in ((16, rel), (17, note or None)):
        calls.cell(row=r, column=col, value=value)
    if verify:
        cell = calls.cell(row=r, column=18, value=verify)
        cell.hyperlink, cell.font = verify, LINK
    else:
        calls.cell(row=r, column=18, value="tvoj zoznam")
    for col in (14, 15, 16, 17, 18):
        c = calls.cell(row=r, column=col)
        c.border, c.alignment = BORDER, WRAP
        if c.font != LINK:
            c.font = BODY
    if src == W:
        calls.cell(row=r, column=2).font = Font(name=ARIAL, size=10, bold=True, color="006100")

last = len(rows) + 1
calls.freeze_panes = "E2"
calls.auto_filter.ref = f"A1:R{last}"
stav = f"N2:N{last}"
calls.conditional_formatting.add(stav, FormulaRule(formula=['N2="Po uzávierke"'],
                                                   fill=PatternFill("solid", fgColor="D9D9D9"),
                                                   font=Font(name=ARIAL, color="7F7F7F")))
calls.conditional_formatting.add(stav, FormulaRule(formula=['N2="Končí do 14 dní"'],
                                                   fill=PatternFill("solid", fgColor="FFC7CE")))
calls.conditional_formatting.add(stav, FormulaRule(formula=[f'N2="{OPEN}"'],
                                                   fill=PatternFill("solid", fgColor="C6EFCE")))
calls.conditional_formatting.add(stav, FormulaRule(formula=[f'N2="{EXP}"'],
                                                   fill=PatternFill("solid", fgColor="DDEBF7")))
calls.conditional_formatting.add(f"P2:P{last}", FormulaRule(formula=['P2="Vysoká"'],
                                                            font=Font(name=ARIAL, bold=True, color="006100")))
calls.conditional_formatting.add(f"A2:M{last}", FormulaRule(formula=['$N2="Po uzávierke"'],
                                                            font=Font(name=ARIAL, color="A6A6A6")))
dv = DataValidation(type="list", formula1='"Vysoká,Stredná,Nízka"', allow_blank=True)
calls.add_data_validation(dv)
dv.add(f"P2:P{last}")

# --- Súhrn ---------------------------------------------------------------------------------
summary["A1"] = "Porovnanie dotačných výziev ČR – grantové poradenstvo"
summary["A1"].font = Font(name=ARIAL, bold=True, size=14)
summary["A3"], summary["B3"] = "Dnes (stav sa prepočíta od tohto dátumu):", TODAY
summary["A3"].font = Font(name=ARIAL, bold=True, size=10)
summary["B3"].font, summary["B3"].fill, summary["B3"].number_format = BLUE, INPUT_FILL, "d. m. yyyy"
summary["B3"].comment = Comment("Vstup: zmeň na dnešný dátum, stĺpce Stav a Dní do uzávierky sa prepočítajú.", "Apoliak7777")

states = ["Po uzávierke", "Končí do 14 dní", OPEN, EXP, CONT]
rels = ["Vysoká", "Stredná", "Nízka"]
top = 5
summary.cell(row=top, column=1, value="Stav \\ Relevancia")
for j, rel in enumerate(rels + ["Spolu"], 2):
    summary.cell(row=top, column=j, value=rel)
for i, st in enumerate(states + ["Spolu"], top + 1):
    summary.cell(row=i, column=1, value=st)
    for j, rel in enumerate(rels, 2):
        col = get_column_letter(j)
        if st == "Spolu":
            summary.cell(row=i, column=j, value=f"=SUM({col}{top + 1}:{col}{i - 1})")
        else:
            summary.cell(row=i, column=j, value=(
                f"=COUNTIFS('Výzvy'!$N$2:$N${last},$A{i},'Výzvy'!$P$2:$P${last},{col}${top})"))
    summary.cell(row=i, column=5, value=f"=SUM(B{i}:D{i})")
end1 = top + len(states) + 1

top2 = end1 + 2
summary.cell(row=top2, column=1, value="Zdroj \\ Stav")
for j, st in enumerate(states + ["Spolu"], 2):
    summary.cell(row=top2, column=j, value=st)
for i, src in enumerate([L, W, "Spolu"], top2 + 1):
    summary.cell(row=i, column=1, value=src)
    for j, st in enumerate(states, 2):
        col = get_column_letter(j)
        if src == "Spolu":
            summary.cell(row=i, column=j, value=f"=SUM({col}{top2 + 1}:{col}{i - 1})")
        else:
            summary.cell(row=i, column=j, value=(
                f"=COUNTIFS('Výzvy'!$B$2:$B${last},$A{i},'Výzvy'!$N$2:$N${last},{col}${top2})"))
    summary.cell(row=i, column=len(states) + 2, value=f"=SUM(B{i}:{get_column_letter(len(states) + 1)}{i})")
end2 = top2 + 3

top3 = end2 + 2
summary.cell(row=top3, column=1, value="Oplatí sa riešiť hneď (vysoká relevancia, ešte sa dá podať)")
summary.cell(row=top3 + 1, column=1, value="Počet:")
summary.cell(row=top3 + 1, column=2, value=(
    f"=COUNTIFS('Výzvy'!$P$2:$P${last},\"Vysoká\",'Výzvy'!$N$2:$N${last},\"<>Po uzávierke\")"))

for row in summary.iter_rows(min_row=top, max_row=top3 + 1):
    for cell in row:
        if cell.value is not None:
            cell.font = BODY
            if cell.row in (top, top2) or cell.column == 1:
                cell.font = Font(name=ARIAL, bold=True, size=10)
            if cell.row in (top, top2):
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
            cell.border = BORDER
summary.cell(row=top3, column=1).border = Border()
summary.column_dimensions["A"].width = 44
for col in "BCDEFG":
    summary.column_dimensions[col].width = 16

# --- Legenda -------------------------------------------------------------------------------
notes = [
    ("Čo je v zošite", ""),
    ("Výzvy", f"{len(R)} výziev: {sum(1 for r in R if r[0] == L)} z tvojho zoznamu + "
              f"{sum(1 for r in R if r[0] == W)} doplnkov z webu, ktoré v zozname chýbali. Zoradené: najprv tie, "
              "na ktoré sa dá podať, podľa relevancie a uzávierky; stĺpec ID = poradie v zdroji."),
    ("Zdroj", "„Tvoj zoznam“ = údaje prevzaté z tvojho zoznamu (neoverené proti poskytovateľom). "
              "„Doplnok (web)“ = nájdené vyhľadávaním 25. 9. 2026; odkaz v stĺpci Overenie."),
    ("Stav k dátumu", "Vzorec z Uzávierky a dátumu v Súhrn!B3: Po uzávierke / Končí do 14 dní / Otvorená; "
                      "bez dátumu sa preberá Status v zdroji (Priebežná, Očakávaná)."),
    ("Relevancia", "Môj odhad pre komerčné grantové poradenstvo v ČR: Vysoká = široký okruh klientov, "
                   "rozumný termín a objem; Nízka = medzinárodné konzorciá, obrana, infraštruktúra, tendre, "
                   "nadácie s malými sumami. Dá sa prepísať (rozbaľovací zoznam)."),
    ("Sumy", "Alokácia a max. dotácia sú text tak, ako v zdroji (Kč aj €); kde mena v zdroji vyzerá chybne, "
             "je to v poznámke."),
    ("Upozornenia", "Stránka optak.gov.cz bola z prostredia nedostupná; doplnky OP TAK, IROP, NRB a SZIF sú "
                    "z výsledkov vyhľadávania – pred ponukou klientovi overiť u poskytovateľa."),
    ("Opravy po overení", "Overené 25. 9. 2026 (výsledky vyhľadávania, oficiálne stránky poskytovateľov); pôvodný "
                          "termín zo zoznamu je v stĺpci Termín – poznámka. Uzavreté, hoci ich zoznam uvádzal ako "
                          "otvorené: AOPK 17. a 18. výzva (2. 6. 2026), OPŽP 104. výzva (30. 6. 2026), NZÚ HOUSEnerg "
                          "bytové domy A–D (8. 11. 2025; nástupca NZÚ 2026+ = riadok 117), OP ST Transformačný úver "
                          "(pozastavený 10. 6. 2026). EIC Accelerator: posledný cut-off 2026 je 4. 11. Výzvu OPZ+ 092 "
                          "(detské skupiny) sa overiť nepodarilo."),
    ("Nepridané (už uzavreté)", "OP TAK Digitální podnik I (18. 2. 2026) a II (17. 4. 2026), Inovace IV "
                                "(20. 2. 2026), Potenciál III (17. 2. 2026), NRB Expanze (pozastavené od "
                                "10. 8. 2026), EDIH II (25. 5. 2026), TA ČR TREND podprogram 2 (v 2026 sa nevyhlási)."),
    ("Sledovať", "Nová výzva OP TAK Úspory energie (plánovaná na 2026, termín neznámy); ďalšie kolo Digitální "
                 "podnik (pre ERP/Odoo najdôležitejšie) – nenašiel som."),
]
legend.column_dimensions["A"].width = 24
legend.column_dimensions["B"].width = 110
for i, (key, text) in enumerate(notes, 1):
    a, b = legend.cell(row=i, column=1, value=key), legend.cell(row=i, column=2, value=text or None)
    a.font = Font(name=ARIAL, bold=True, size=12 if i == 1 else 10)
    b.font, b.alignment = BODY, WRAP

out = str(Path(__file__).with_name("vyzvy_CZ_porovnanie_2026-09-25.xlsx"))
wb.properties.creator = "Apoliak7777"
wb.properties.title = "Porovnanie dotačných výziev ČR"
wb.save(out)
print(out, len(R), "rows")
